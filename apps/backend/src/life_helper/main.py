"""Local foundation API with a separate, preserved anonymous experiment area."""

import hashlib
import hmac
import io
import os
import re
import secrets
import uuid
import warnings
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Annotated, Literal
from urllib.parse import urlsplit

from docx import Document
from fastapi import (
    Depends,
    FastAPI,
    File,
    Form,
    Header,
    HTTPException,
    Request,
    Response,
    UploadFile,
)
from fastapi.responses import FileResponse, JSONResponse
from PIL import Image, ImageOps, UnidentifiedImageError
from pydantic import BaseModel, Field, model_validator
from sqlalchemy import JSON, Column, String, create_engine, select
from sqlalchemy.orm import Session, declarative_base

from life_helper.foundation import attach_foundation
from life_helper.request_limits import RequestBodyLimit

Base = declarative_base()
MAX_UPLOAD = 10 * 1024 * 1024
MAX_PIXELS = 20_000_000
COOKIE = "life_workspace"
COMPONENTS = ("images.resize", "documents.generate", "images.remove-background", "video.generate")


class Workspace(Base):
    __tablename__ = "preview_workspaces"
    id = Column(String, primary_key=True)
    token_hash = Column(String, unique=True, nullable=False)
    expires_at = Column(String, nullable=False)


class Job(Base):
    __tablename__ = "jobs"
    id = Column(String, primary_key=True)
    workspace_id = Column(String, index=True, nullable=False)
    kind = Column(String, nullable=False)
    status = Column(String, nullable=False)
    created_at = Column(String, nullable=False)
    title = Column(String, nullable=False)
    result_path = Column(String, nullable=False)
    details = Column(JSON, nullable=False)


class Binding(Base):
    __tablename__ = "component_bindings"
    component = Column(String, primary_key=True)
    config = Column(JSON, nullable=False)


class RouteConfig(BaseModel):
    component: Literal[
        "images.resize", "documents.generate", "images.remove-background", "video.generate"
    ]
    mode: Literal["local", "mock", "api_draft"]
    provider: str = Field(default="", max_length=80)
    model: str = Field(default="", max_length=120)
    base_url: str = Field(default="", max_length=300)
    secret_env: str = Field(default="", max_length=100)

    @model_validator(mode="after")
    def valid_binding(self):
        local = self.component == "images.resize" and self.mode == "local"
        mock = self.component == "documents.generate" and self.mode == "mock"
        if self.mode != "api_draft" and not (local or mock):
            raise ValueError("该组件暂不支持此执行模式")
        if self.mode != "api_draft" and any(
            (self.provider, self.model, self.base_url, self.secret_env)
        ):
            raise ValueError("本地或模拟模式请清空外部接口字段")
        if self.secret_env and not re.fullmatch(r"[A-Z][A-Z0-9_]{2,99}", self.secret_env):
            raise ValueError("仅填写密钥环境变量名，例如 TEXT_PROVIDER_API_KEY")
        if self.base_url:
            url = urlsplit(self.base_url)
            if (
                url.scheme != "https"
                or not url.hostname
                or url.username
                or url.password
                or url.query
                or url.fragment
            ):
                raise ValueError("接口地址必须为不含凭据和查询参数的HTTPS地址")
        return self


class DocumentInput(BaseModel):
    company: str = Field(min_length=1, max_length=80)
    product: str = Field(min_length=1, max_length=200)
    audience: str = Field(default="客户", max_length=80)
    template: Literal["product", "notice"] = "product"


class DocumentEdit(BaseModel):
    text: str = Field(min_length=1, max_length=20000)


def now():
    return datetime.now(UTC).isoformat()


def default_bindings():
    return [
        RouteConfig(component="images.resize", mode="local"),
        RouteConfig(component="documents.generate", mode="mock"),
        RouteConfig(component="images.remove-background", mode="api_draft"),
        RouteConfig(component="video.generate", mode="api_draft"),
    ]


def create_app(data_dir: Path | None = None, admin_token: str | None = None):
    root = (data_dir or Path(os.getenv("LIFE_DATA_DIR", "runtime"))).resolve()
    root.mkdir(parents=True, exist_ok=True)
    results = root / "results"
    results.mkdir(exist_ok=True)
    engine = create_engine(
        f"sqlite:///{root / 'preview.sqlite3'}", connect_args={"check_same_thread": False}
    )
    Base.metadata.create_all(engine)
    owner_token = admin_token if admin_token is not None else os.getenv("LIFE_ADMIN_TOKEN", "")
    origins = set(
        os.getenv("LIFE_ALLOWED_ORIGINS", "http://localhost:3000,http://127.0.0.1:3000").split(",")
    )
    app = FastAPI(title="生活小助手 · 本机开发预览", version="0.1.0")
    app.state.engine = engine

    with Session(engine) as db:
        for route in default_bindings():
            if db.get(Binding, route.component) is None:
                db.add(Binding(component=route.component, config=route.model_dump()))
        db.commit()

    @app.middleware("http")
    async def guard_request(request: Request, call_next):
        if request.method not in {"GET", "HEAD", "OPTIONS"}:
            origin = request.headers.get("origin")
            if (origin is not None and origin not in origins) or request.headers.get(
                "sec-fetch-site"
            ) == "cross-site":
                return JSONResponse(status_code=403, content={"detail": "不允许此请求来源"})
            length = request.headers.get("content-length", "0")
            if not length.isdigit() or int(length) > MAX_UPLOAD + 1024 * 1024:
                return JSONResponse(status_code=413, content={"detail": "上传文件超过10MB限制"})
        response = await call_next(request)
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["Cache-Control"] = "no-store"
        return response

    def database():
        with Session(engine) as db:
            yield db

    Database = Annotated[Session, Depends(database)]

    def current_workspace(request: Request, db: Database):
        token = request.cookies.get(COOKIE, "")
        if not token or len(token) > 200:
            raise HTTPException(401, "请先开启本机工作空间")
        digest = hashlib.sha256(token.encode()).hexdigest()
        workspace = db.scalar(select(Workspace).where(Workspace.token_hash == digest))
        if workspace is None or workspace.expires_at < now():
            raise HTTPException(401, "工作空间已过期，请重新开启")
        return workspace.id

    CurrentWorkspace = Annotated[str, Depends(current_workspace)]

    def require_owner(authorization: Annotated[str | None, Header()] = None):
        if len(owner_token) < 24:
            raise HTTPException(503, "站长配置未开启：请先设置至少24位的 LIFE_ADMIN_TOKEN")
        supplied = (authorization or "").removeprefix("Bearer ")
        if not hmac.compare_digest(supplied, owner_token):
            raise HTTPException(403, "站长口令不正确")

    def job_for(db: Session, job_id: str, workspace_id: str):
        job = db.get(Job, job_id)
        if job is None or job.workspace_id != workspace_id:
            raise HTTPException(404, "任务不存在")
        return job

    def serialize(job: Job):
        return {
            "id": job.id,
            "kind": job.kind,
            "status": job.status,
            "created_at": job.created_at,
            "title": job.title,
            "details": job.details,
            "download_url": f"/api/jobs/{job.id}/download",
        }

    @app.get("/api/health")
    def health():
        return {"status": "ok", "mode": "local_preview", "billable_calls_enabled": False}

    @app.post("/api/session")
    def start_session(request: Request, response: Response, db: Database):
        try:
            workspace_id = current_workspace(request, db)
            return {"workspace_id": workspace_id, "mode": "anonymous_preview"}
        except HTTPException:
            pass
        token = secrets.token_urlsafe(32)
        workspace_id = str(uuid.uuid4())
        db.add(
            Workspace(
                id=workspace_id,
                token_hash=hashlib.sha256(token.encode()).hexdigest(),
                expires_at=(datetime.now(UTC) + timedelta(days=7)).isoformat(),
            )
        )
        db.commit()
        response.set_cookie(
            COOKIE,
            token,
            httponly=True,
            samesite="lax",
            max_age=604800,
            secure=os.getenv("LIFE_COOKIE_SECURE", "false").lower() == "true",
        )
        return {"workspace_id": workspace_id, "mode": "anonymous_preview"}

    @app.get("/api/jobs")
    def list_jobs(db: Database, workspace_id: CurrentWorkspace):
        jobs = db.scalars(
            select(Job)
            .where(Job.workspace_id == workspace_id)
            .order_by(Job.created_at.desc())
            .limit(50)
        ).all()
        return {"jobs": [serialize(job) for job in jobs]}

    @app.post("/api/images/resize")
    def resize_image(
        db: Database,
        workspace_id: CurrentWorkspace,
        file: Annotated[UploadFile, File()],
        width: Annotated[int, Form(ge=32, le=4096)] = 1200,
        quality: Annotated[int, Form(ge=30, le=95)] = 85,
        output_format: Annotated[Literal["jpeg", "png", "webp"], Form()] = "jpeg",
    ):
        binding = db.get(Binding, "images.resize")
        if binding.config["mode"] != "local":
            raise HTTPException(409, "该组件已改为接口草稿，尚未接通，请站长切回本地模式")
        raw = file.file.read(MAX_UPLOAD + 1)
        if len(raw) > MAX_UPLOAD:
            raise HTTPException(413, "图片超过10MB限制")
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(io.BytesIO(raw)) as source:
                    if source.format not in {"PNG", "JPEG", "WEBP"}:
                        raise HTTPException(422, "请上传PNG、JPG或WebP图片")
                    if source.width * source.height > MAX_PIXELS:
                        raise HTTPException(422, "图片像素超过2000万限制")
                    image = ImageOps.exif_transpose(source)
                    original_size = image.size
                    target_width = min(width, image.width)
                    target_height = max(1, round(image.height * target_width / image.width))
                    if target_width * target_height > MAX_PIXELS:
                        raise HTTPException(422, "输出像素超过限制")
                    image = image.resize((target_width, target_height), Image.Resampling.LANCZOS)
                    if output_format == "jpeg":
                        rgba = image.convert("RGBA")
                        image = Image.new("RGB", image.size, "white")
                        image.paste(rgba, mask=rgba.getchannel("A"))
                    elif image.mode not in {"RGB", "RGBA"}:
                        image = image.convert("RGBA")
                    job_id = str(uuid.uuid4())
                    path = results / f"{job_id}.{output_format}"
                    image.info.clear()
                    image.save(path, format=output_format.upper(), quality=quality)
        except (
            UnidentifiedImageError,
            OSError,
            ValueError,
            Image.DecompressionBombError,
            Image.DecompressionBombWarning,
        ) as error:
            raise HTTPException(422, "图片无法读取，请换一张有效的PNG、JPG或WebP") from error
        details = {
            "input_bytes": len(raw),
            "output_bytes": path.stat().st_size,
            "original_width": original_size[0],
            "original_height": original_size[1],
            "width": image.width,
            "height": image.height,
            "format": output_format,
            "mode": "local",
            "api_cost": 0,
        }
        title = (file.filename or "图片").replace("\\", "/").rsplit("/", 1)[-1][:160]
        job = Job(
            id=job_id,
            workspace_id=workspace_id,
            kind="image",
            status="succeeded",
            created_at=now(),
            title=title,
            result_path=str(path),
            details=details,
        )
        db.add(job)
        db.commit()
        return serialize(job)

    @app.post("/api/documents/preview")
    def document_preview(payload: DocumentInput, db: Database, workspace_id: CurrentWorkspace):
        binding = db.get(Binding, "documents.generate")
        if binding.config["mode"] != "mock":
            raise HTTPException(409, "文本API尚未接通，当前配置仅保存为草稿")
        company, product = payload.company.strip(), payload.product.strip()
        if not company or not product:
            raise HTTPException(422, "请填写企业名称和产品或事项")
        title = "产品介绍" if payload.template == "product" else "客户通知"
        text = (
            f"{company}｜{title}\n\n致{payload.audience.strip() or '客户'}：\n\n"
            f"本次需要介绍或说明的内容：{product}。\n\n"
            "请补充准确的产品特点、适用场景、价格或时间、联系方式。"
            "发布前请核实所有事实，并删除不适用的提示。\n\n"
            "【模板预览：未调用AI，未产生模型费用】"
        )
        job_id = str(uuid.uuid4())
        path = results / f"{job_id}.docx"
        write_docx(path, title, text)
        job = Job(
            id=job_id,
            workspace_id=workspace_id,
            kind="document",
            status="succeeded",
            created_at=now(),
            title=f"{company} · {title}",
            result_path=str(path),
            details={"mode": "mock", "text": text, "api_cost": 0},
        )
        db.add(job)
        db.commit()
        return serialize(job)

    @app.post("/api/jobs/{job_id}/document")
    def edit_document(
        job_id: str, payload: DocumentEdit, db: Database, workspace_id: CurrentWorkspace
    ):
        job = job_for(db, job_id, workspace_id)
        if job.kind != "document":
            raise HTTPException(422, "该任务不是文档")
        if not payload.text.strip():
            raise HTTPException(422, "文档内容不能为空")
        write_docx(Path(job.result_path), job.title, payload.text)
        job.details = {**job.details, "text": payload.text}
        db.commit()
        return serialize(job)

    @app.get("/api/jobs/{job_id}/download")
    def download(job_id: str, db: Database, workspace_id: CurrentWorkspace):
        job = job_for(db, job_id, workspace_id)
        path = Path(job.result_path)
        if not path.is_file():
            raise HTTPException(404, "结果文件已不存在")
        return FileResponse(path, filename=f"lifeisgood-{job.id[:8]}{path.suffix}")

    @app.get("/api/admin/bindings", dependencies=[Depends(require_owner)])
    def get_bindings(db: Database):
        return {
            "bindings": [row.config for row in db.scalars(select(Binding)).all()],
            "billable_calls_enabled": False,
        }

    @app.put("/api/admin/bindings", dependencies=[Depends(require_owner)])
    def update_binding(payload: RouteConfig, db: Database):
        row = db.get(Binding, payload.component)
        row.config = payload.model_dump()
        db.commit()
        return {"binding": row.config, "billable_calls_enabled": False}

    attach_foundation(app, root)
    app.add_middleware(
        RequestBodyLimit,
        max_bytes=MAX_UPLOAD + 1024 * 1024,
        authenticate=app.state.foundation_authenticate,
        origins=origins,
    )
    return app


def write_docx(path: Path, title: str, text: str):
    document = Document()
    document.add_heading(title, 0)
    for paragraph in text.split("\n"):
        document.add_paragraph(paragraph)
    document.add_paragraph("生活小助手 · 模板预览，未调用AI。请核实后使用。")
    temporary = path.with_suffix(".tmp")
    document.save(temporary)
    temporary.replace(path)


app = create_app()
